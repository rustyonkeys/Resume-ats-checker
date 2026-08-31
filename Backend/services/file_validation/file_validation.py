from pathlib import Path
from docx import Document
import fitz



def validate_text_quality(
        text: str, 
        min_words: int=40 , 
        min_chars: int=100) -> bool:
    """To check if the extracted text meets the minimum content thresholds. """
    clean_text = text.strip()

    #Splits by whitespace to find geniune words
    words = clean_text.split()

    if len(words)<(min_words):
        raise ValueError(
            f"Resume contains insufficient text ({len(words)} words found). "
            f"Minimum required is {min_words} words."
        )
    
    if len(clean_text)<min_chars:
        raise ValueError(
            f"Resume contains insufficient text ({len(clean_text)} characters found). "
            f"Minimum required is {min_chars} characters."
        )
    return True


def file_validation(
        file_path: str,
        allowed_file_extensions: set[str],
        max_mb: int,
        min_words: int=40 , 
        min_chars: int=100) ->dict:
# File extension and size validation

    #Validate an uploaded resume file.

    # Raises:
    #     ValueError: When the file fails validation.

    # Returns:
    #     dict: Basic validation information.
        path = Path(file_path)

        if not path.exists() or not path.is_file():
            raise ValueError(f"Uploaded file cannot be found")
        
        file_size_bytes = path.stat().st_size
        max_size_bytes = max_mb * 1024 * 1024
        extension = path.suffix.lower()

        # validate extension
        if extension not in allowed_file_extensions:
            raise ValueError(
                f"Unsupported file extension"
                f"Allowed extensions are: {', '.join(sorted(allowed_file_extensions))}")

        # validate empty file
        if file_size_bytes == 0:
            raise ValueError(f"File is empty")
        
        # validate file size
        if file_size_bytes > max_size_bytes:
            raise ValueError(f"File size exceeds the maximum allowed size of {max_mb} MB")
        
        page_count = None
        contains_text = None

        full_resume_text = ""

        # Deep pdf content validation
        if extension == '.pdf':
            try:
                with fitz.open(file_path) as document:
                    if document.page_count == 0:
                        raise ValueError(f"PDF is corrupted or contains no pages")
                    
                    if document.needs_pass:
                        raise ValueError(f"Password-protected pdfs are not supported")
                    
                    page_count = document.page_count

                    pdf_text = []
                    for page in document:
                        page_text = page.get_text().strip()

                        if page_text:
                            pdf_text.append(page_text)

                    full_resume_text = "\n".join(pdf_text)

                    contains_text = bool(full_resume_text)

                    if not contains_text:
                        raise ValueError(
                            f"PDF contains no extractable text"
                            f"It may be a scanned or image only PDF, which is not supported.")
                    
                    
                    validate_text_quality(
                        full_resume_text,
                        min_words = min_words,
                        min_chars = min_chars,
                    )
                        

                    
    # ValueError created by you
    # return the same error
            except ValueError:
                raise
    # Unexpected library/system error
    # convert it into a clearer ValueError

            except Exception as exc:
                raise ValueError(f"Pdf file is corrupted or unreadable: {str(exc)}") from exc
            
        elif extension == ".docx":
            try:
                document = Document(file_path)
            #     extracted_text = "\n".join(
            #     paragraph.text.strip()
            #     for paragraph in document.paragraphs
            #     if paragraph.text.strip()
            # )

                text_parts = []

                for paragraph in document.paragraphs:
                    paragraph_text = paragraph.text.strip()

                    if paragraph_text:
                        text_parts.append(paragraph_text)    

                full_resume_text = "\n".join(text_parts)
                contains_text = bool(full_resume_text)

                if not contains_text:
                    raise ValueError(
                        "DOCX contains no extractable text")
                    
                validate_text_quality(
                    full_resume_text,
                    min_words = min_words,
                    min_chars = min_chars,
                )

            except ValueError:
                raise

            except Exception as exc:
                raise ValueError(
                f"DOCX file is corrupted or unreadable: {exc}"
            ) from exc
            
        return {
            "valid": True,
            "extension": extension,
            "size_bytes": file_size_bytes,
            "page_count": page_count,
            "contains_extractable_text": contains_text,
            "full_resume_text": full_resume_text
        }




